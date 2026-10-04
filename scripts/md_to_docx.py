#!/usr/bin/env python3
"""Convert a Markdown file to a styled .docx for Google Docs import.

The point is to skip hand-formatting: the output uses real Word heading styles
and real tables, so Google Docs converts it into a properly structured
document (and a navigable outline) rather than one giant blob of preformatted
text.

Usage:
    python scripts/md_to_docx.py SOLUTION_DETAILS.md
    python scripts/md_to_docx.py SOLUTION_DETAILS.md -o /tmp/out.docx
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    import docx  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Pt, RGBColor  # type: ignore
except ImportError:  # pragma: no cover
    sys.exit("python-docx is required:  pip install python-docx")

MONO = "Consolas"
BODY = "Calibri"
HEADING_COLOURS = {
    1: RGBColor(0x1A, 0x1A, 0x1A),
    2: RGBColor(0x0B, 0x4F, 0x6C),
    3: RGBColor(0x1A, 0x1A, 0x1A),
    4: RGBColor(0x44, 0x44, 0x44),
}
CODE_BG = "F4F5F7"
INLINE_BG = "EEF1F4"


# --------------------------------------------------------------- inline text

def add_inline(paragraph, text: str, *, base_bold: bool = False,
               base_size: int | None = None, base_font: str = BODY,
               base_colour: RGBColor | None = None) -> None:
    """Render **bold**, *italic* and `code` spans as separate runs."""
    # Split on the three inline markers, keeping the delimiters.
    pattern = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`)")
    for part in pattern.split(text):
        if not part:
            continue
        bold, italic, mono = base_bold, False, False
        body = part
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            bold, body = True, part[2:-2]
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            mono, body = True, part[1:-1]
        elif (part.startswith("*") and part.endswith("*")
              and len(part) > 2 and not part.startswith("**")):
            italic, body = True, part[1:-1]

        run = paragraph.add_run(body)
        run.bold = bold
        run.italic = italic
        run.font.name = MONO if mono else base_font
        if base_size:
            run.font.size = Pt(base_size - 1 if mono else base_size)
        if base_colour is not None:
            run.font.color.rgb = base_colour
        if mono:
            run.font.color.rgb = RGBColor(0x0B, 0x3D, 0x2E)
            run.font.size = Pt(base_size - 1 if base_size else 9)


def shade(element, hex_fill: str) -> None:
    """Apply a background fill to a paragraph or table cell."""
    from docx.oxml.ns import qn  # type: ignore
    from docx.oxml import OxmlElement  # type: ignore

    pr = element._p.get_or_add_pPr() if hasattr(element, "_p") \
        else element._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


# ------------------------------------------------------------------ elements

def add_code_block(doc, lines: list[str]) -> None:
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.left_indent = Pt(12)
    pf.space_before = Pt(6)
    pf.space_after = Pt(10)
    pf.line_spacing = 1.0
    shade(para, CODE_BG)
    for i, line in enumerate(lines):
        run = para.add_run(line)
        run.font.name = MONO
        run.font.size = Pt(9)
        if i < len(lines) - 1:
            run.add_break()
    return para


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def add_table(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    try:
        table.style = doc.styles["Table Grid"]
    except KeyError:
        pass
    table.autofit = True

    for cell, text in zip(table.rows[0].cells, header):
        cell.text = ""
        para = cell.paragraphs[0]
        add_inline(para, text, base_bold=True, base_size=10)
        shade(cell, "E8EDF2")

    for row in body:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            add_inline(cell.paragraphs[0], text, base_size=10)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_bullet(doc, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else f"List Bullet {level + 1}"
    try:
        para = doc.add_paragraph(style=style)
    except KeyError:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Pt(18 + 18 * level)
    para.paragraph_format.space_after = Pt(3)
    add_inline(para, text, base_size=11)


# ----------------------------------------------------------------- converter

def convert(md_path: Path, out_path: Path) -> Path:
    doc = docx.Document()

    normal = doc.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i, n = 0, len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # fenced code block
        if stripped.startswith("```"):
            i += 1
            block: list[str] = []
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, block)
            continue

        # horizontal rule
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            from docx.oxml import OxmlElement  # type: ignore
            from docx.oxml.ns import qn  # type: ignore
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "CCCCCC")
            pbdr.append(bottom)
            para._p.get_or_add_pPr().append(pbdr)
            i += 1
            continue

        # table: a pipe row followed by a separator row
        if ("|" in stripped and i + 1 < n and is_separator(lines[i + 1])):
            rows = [split_row(stripped)]
            i += 2
            while i < n and "|" in lines[i] and lines[i].strip():
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            para = doc.add_heading(level=min(level, 4))
            para.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            para.paragraph_format.space_after = Pt(5)
            run = para.add_run(text)
            run.font.name = BODY
            run.font.size = Pt({1: 20, 2: 15, 3: 12.5, 4: 11}[min(level, 4)])
            run.font.color.rgb = HEADING_COLOURS.get(min(level, 4),
                                                       RGBColor(0, 0, 0))
            run.bold = True
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            quote: list[str] = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(16)
            para.paragraph_format.space_before = Pt(6)
            shade(para, "FFF8E1")
            add_inline(para, " ".join(q for q in quote if q.strip()).strip(),
                       base_size=10)
            continue

        # bullet list
        m = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if m:
            level = len(m.group(1)) // 2
            add_bullet(doc, m.group(2), level)
            i += 1
            continue

        # numbered list
        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            try:
                para = doc.add_paragraph(style="List Number")
            except KeyError:
                para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(3)
            add_inline(para, m.group(1), base_size=11)
            i += 1
            continue

        # blank
        if not stripped:
            i += 1
            continue

        # paragraph: consume until a blank line or a new block starts
        buf = [stripped]
        i += 1
        while i < n and lines[i].strip() and not re.match(
                r"^\s*(#{1,6}\s|[-*+]\s|\d+\.\s|>|\||```)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        para = doc.add_paragraph()
        add_inline(para, " ".join(buf), base_size=11)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("markdown", type=Path)
    ap.add_argument("-o", "--out", type=Path, default=None)
    args = ap.parse_args()

    if not args.markdown.is_file():
        sys.exit(f"no such file: {args.markdown}")

    out = args.out or args.markdown.with_suffix(".docx")
    convert(args.markdown, out)
    size_kb = out.stat().st_size / 1024
    print(f"wrote {out}  ({size_kb:.0f} KB)")


if __name__ == "__main__":
    main()
